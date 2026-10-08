"""Independently reopen browser downloads with Pillow and openpyxl."""
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from openpyxl import load_workbook

actual = Path(__file__).resolve().parent.parent / "tests" / "actual"
expected_images = {
    "compressed.jpg": ("JPEG", (1100, 720)),
    "compressed-strong.jpg": ("JPEG", (1100, 720)),
    "compressed.png": ("PNG", (1100, 720)),
    "resize-width.jpg": ("JPEG", (550, 360)),
    "resize-height.jpg": ("JPEG", (550, 360)),
    "resize-exact.jpg": ("JPEG", (400, 300)),
    "transparent-blue.jpg": ("JPEG", (80, 60)),
    "transparent.webp": ("WEBP", (80, 60)),
    "converted.png": ("PNG", (1100, 720)),
    "target-50kb.webp": ("WEBP", (1100, 720)),
    "target-resized.webp": ("WEBP", (704, 461)),
    "webp-to-png.png": ("PNG", (80, 60)),
    "webp-to-jpg.jpg": ("JPEG", (80, 60)),
    "oriented-output.jpg": ("JPEG", (40, 80)),
    "tiny-output.jpg": ("JPEG", (1, 1)),
    "metadata-removed.jpg": ("JPEG", (1100, 720)),
}
for name, (fmt, dimensions) in expected_images.items():
    with Image.open(actual / name) as image:
        image.load()
        assert image.format == fmt, (name, image.format)
        assert image.size == dimensions, (name, image.size)
        if name == "transparent.webp":
            assert image.convert("RGBA").getpixel((0, 0))[3] == 0
        if name == "metadata-removed.jpg":
            exif = image.getexif()
            assert 34853 not in exif and 271 not in exif and 272 not in exif

with ZipFile(actual / "batch.zip") as archive:
    assert archive.testzip() is None
    assert len(archive.namelist()) == 2
    for name in archive.namelist():
        with Image.open(archive.open(name)) as image:
            image.load()
            assert image.format == "JPEG"

with ZipFile(actual / "optimized.xlsx") as archive:
    assert archive.testzip() is None
workbook = load_workbook(actual / "optimized.xlsx", read_only=False, data_only=False)
assert "한글 시트" in workbook.sheetnames
assert workbook["한글 시트"]["A1"].value == "반복 자료 1"
assert workbook["한글 시트"]["P1"].value == "=SUM(1,2)"
workbook.close()
with ZipFile(actual / "optimized-with-dirs.xlsx") as archive:
    assert archive.testzip() is None
    assert "xl/" in archive.namelist()
    assert "xl/worksheets/" in archive.namelist()
directory_book = load_workbook(actual / "optimized-with-dirs.xlsx")
assert directory_book["한글 시트"]["P1"].value == "=SUM(1,2)"
directory_book.close()
complex_book = load_workbook(actual / "complex-optimized.xlsx", read_only=False, data_only=False)
assert complex_book["숨김 시트"].sheet_state == "hidden"
assert complex_book["숨김 시트"]["A1"].value == "보존할 데이터"
assert complex_book["공개 시트"]["D2"].value == "=SUM(B2:B11)"
assert "D4:E4" in [str(r) for r in complex_book["공개 시트"].merged_cells.ranges]
assert len(complex_book["공개 시트"]._charts) == 1
assert len(complex_book["공개 시트"]._images) == 1
complex_book.close()
print(f"Verified {len(expected_images)} images, batch ZIP, and 3 XLSX files with Pillow/openpyxl")
