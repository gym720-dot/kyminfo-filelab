"""수동 브라우저 검증용 예제 파일. Python Pillow/openpyxl이 필요합니다."""
from pathlib import Path
from PIL import Image, ImageDraw
from openpyxl import Workbook

out = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
out.mkdir(parents=True, exist_ok=True)
image = Image.new("RGB", (1100, 720), "#dfe8ff")
draw = ImageDraw.Draw(image)
for i in range(180):
    draw.rectangle((i * 7 % 1100, i * 19 % 720, (i * 7 % 1100) + 90, (i * 19 % 720) + 45), fill=(i * 23 % 255, i * 47 % 255, i * 89 % 255))
image.save(out / "sample.jpg", quality=95)
image.save(out / "sample.png")

wb = Workbook()
ws = wb.active
ws.title = "한글 시트"
for row in range(1, 260):
    for col in range(1, 15):
        ws.cell(row, col, "반복 자료 " + str(row % 9))
ws["P1"] = "=SUM(1,2)"
wb.save(out / "sample.xlsx")
print(out)
